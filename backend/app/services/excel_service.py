"""Excel (.xlsx) generation with OpenPyXL, preserving Unicode & formatting."""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.utils.exceptions import ExcelExportError
from app.utils.logger import logger

_HEADER_FILL = PatternFill("solid", fgColor="4F81BD")
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
_CELL_FONT = Font(size=11)
_THIN = Side(style="thin", color="B0B0B0")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

_MAX_COL_WIDTH = 60
_MIN_COL_WIDTH = 8


def _looks_numeric(value: str) -> bool:
    """Return True if ``value`` is a plain ASCII number safe to store as num.

    Devanagari digits are converted to ASCII upstream (see
    :mod:`app.utils.numeral_converter`), so by the time a cell reaches here a
    numeric value is already ASCII. Mixed text (e.g. ``विद्यालय नं. 12``) and
    non-numeric strings stay text.
    """
    v = value.strip().replace(",", "")
    if not v or not v.isascii():
        return False
    try:
        float(v)
        return True
    except ValueError:
        return False


class ExcelService:
    """Convert a row-major string table into a formatted workbook."""

    def build_workbook(
        self,
        table: list[list[str]],
        sheet_name: str = "Sheet1",
        header: bool = True,
    ) -> bytes:
        """Return the ``.xlsx`` file as bytes.

        - Unicode (Hindi) text is preserved (openpyxl writes UTF-8 XML).
        - Numeric-looking ASCII values are written as numbers.
        - Header row is bold with fill; all cells get borders, wrapping and
          alignment; the first row is frozen and column widths auto-fit.
        """
        if not table:
            raise ExcelExportError("Cannot export an empty table.")

        try:
            wb = Workbook()
            ws = wb.active
            ws.title = (sheet_name or "Sheet1")[:31]

            n_cols = max(len(r) for r in table)
            col_widths = [float(_MIN_COL_WIDTH)] * n_cols

            for r_idx, row in enumerate(table, start=1):
                is_header = header and r_idx == 1
                for c_idx in range(n_cols):
                    raw = row[c_idx] if c_idx < len(row) else ""
                    value: object = raw
                    cell = ws.cell(row=r_idx, column=c_idx + 1)
                    if not is_header and _looks_numeric(raw):
                        num = float(raw.replace(",", ""))
                        value = int(num) if num.is_integer() else num
                    cell.value = value
                    cell.border = _BORDER
                    if is_header:
                        cell.font = _HEADER_FONT
                        cell.fill = _HEADER_FILL
                        cell.alignment = _CENTER
                    else:
                        cell.font = _CELL_FONT
                        cell.alignment = (
                            _CENTER if _looks_numeric(raw) else _LEFT
                        )
                    # Track widest content per column (longest text line).
                    longest = max(
                        (len(part) for part in str(raw).split("\n")),
                        default=0,
                    )
                    col_widths[c_idx] = max(col_widths[c_idx], longest + 2)

            for c_idx in range(n_cols):
                letter = get_column_letter(c_idx + 1)
                ws.column_dimensions[letter].width = min(
                    col_widths[c_idx], _MAX_COL_WIDTH
                )

            if header and len(table) >= 1:
                ws.freeze_panes = "A2"

            buffer = io.BytesIO()
            wb.save(buffer)
            data = buffer.getvalue()
            logger.info(
                "Generated workbook: {} rows x {} cols ({} bytes)",
                len(table), n_cols, len(data),
            )
            return data
        except ExcelExportError:
            raise
        except Exception as exc:
            logger.exception("Excel generation failed")
            raise ExcelExportError(f"Failed to build Excel file: {exc}") from exc
