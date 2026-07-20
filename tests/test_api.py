"""API-level tests for upload, process and export (OCR mocked)."""
from __future__ import annotations

import io

import openpyxl
from app.dependencies import get_document_service
from app.main import app
from app.models.schemas import PageResult, ProcessResponse


class _FakeDocumentService:
    """Stand-in that returns a deterministic table without running OCR."""

    def process(self, path, file_id, lang=None):
        table = [["A", "B"], ["1", "राम"]]
        return ProcessResponse(
            file_id=file_id,
            pages=[PageResult(page=1, rows=table, n_rows=2, n_cols=2,
                              confidence=0.9)],
            table=table,
            n_rows=2,
            n_cols=2,
            confidence=0.9,
            processing_time=0.01,
        )


def test_upload_then_process(client, png_bytes):
    res = client.post(
        "/api/upload",
        files={"file": ("table.png", png_bytes, "image/png")},
    )
    assert res.status_code == 200
    file_id = res.json()["file_id"]
    assert res.json()["kind"] == "image"

    app.dependency_overrides[get_document_service] = lambda: _FakeDocumentService()
    try:
        res = client.post("/api/process", json={"file_id": file_id})
        assert res.status_code == 200
        body = res.json()
        assert body["table"] == [["A", "B"], ["1", "राम"]]
        assert body["n_rows"] == 2
    finally:
        app.dependency_overrides.clear()


def test_upload_rejects_bad_type(client):
    res = client.post(
        "/api/upload",
        files={"file": ("bad.gif", b"GIF89a....", "image/gif")},
    )
    assert res.status_code == 400
    assert res.json()["error"] == "invalid_file"


def test_process_missing_file_returns_404(client):
    res = client.post("/api/process", json={"file_id": "does-not-exist"})
    assert res.status_code == 404
    assert res.json()["error"] == "file_not_found"


def test_export_returns_xlsx(client, sample_table):
    res = client.post(
        "/api/export",
        json={"table": sample_table, "filename": "out", "header": True},
    )
    assert res.status_code == 200
    assert "spreadsheetml" in res.headers["content-type"]
    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    rows = list(wb.active.iter_rows(values_only=True))
    assert rows[0] == ("क्रम", "नाम", "राशि")


def test_export_empty_table_errors(client):
    res = client.post("/api/export", json={"table": [], "filename": "x"})
    assert res.status_code == 500
    assert res.json()["error"] == "excel_export_failed"


def test_export_converts_devanagari_numerals(client):
    table = [["क्रम", "नाम", "राशि"], ["१", "राम", "१५००"], ["२", "John", "२३०"]]
    res = client.post(
        "/api/export",
        json={"table": table, "filename": "num", "header": True},
    )
    assert res.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    rows = list(wb.active.iter_rows(values_only=True))
    # Header preserved; Devanagari digits become real numeric cells.
    assert rows[0] == ("क्रम", "नाम", "राशि")
    assert rows[1] == (1, "राम", 1500)
    assert rows[2] == (2, "John", 230)
