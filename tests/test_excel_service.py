"""Tests for the Excel export service."""
from __future__ import annotations

import io

import openpyxl
import pytest
from app.services.excel_service import ExcelService, _looks_numeric
from app.utils.exceptions import ExcelExportError


def test_looks_numeric():
    assert _looks_numeric("1500")
    assert _looks_numeric("1,500")
    assert _looks_numeric("12.5")
    assert not _looks_numeric("१२३")  # Devanagari digits are not ASCII
    assert not _looks_numeric("abc")
    assert not _looks_numeric("")


def test_build_workbook_preserves_unicode_and_numbers(sample_table):
    data = ExcelService().build_workbook(sample_table, header=True)
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == ("क्रम", "नाम", "राशि")
    assert rows[1] == (1, "राम कुमार", 1500)  # numeric coercion
    assert ws["A1"].font.bold is True
    assert ws.freeze_panes == "A2"


def test_build_workbook_empty_raises():
    with pytest.raises(ExcelExportError):
        ExcelService().build_workbook([])


def test_build_workbook_ragged_rows_are_padded():
    data = ExcelService().build_workbook([["a", "b", "c"], ["x"]], header=False)
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb.active
    assert ws.max_column == 3
